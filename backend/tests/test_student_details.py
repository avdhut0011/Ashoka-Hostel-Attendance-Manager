from datetime import datetime, timedelta
from io import BytesIO
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend.app.main import app
from backend.app.core.database import TargetSessionLocal
from backend.app.core.security import create_access_token
from backend.app.models import Student, StudentCurrentStatus, User, RawPunch, ProcessedPunch
from backend.app.models.attendance import AttendanceOverride
from backend.app.services.punch_processor import apply_override
from backend.app.services.change_events import change_marker
from backend.app.services import student_details as service

NOW = datetime(2026, 10, 10, 12)


def headers(db, username='admin'):
    user = db.query(User).filter_by(Username=username).one()
    return {'Authorization': 'Bearer ' + create_access_token({'sub': user.Username, 'id': user.SystemUserId, 'ver': 0})}


def student(db, uid, hostel=1, status='OUT', when=None, registration='ACTIVE'):
    row = Student(UserId=uid, StudentName=uid, HostelId=hostel, RoomNo='2', RegistrationStatus=registration)
    db.add(row)
    db.flush()
    db.add(StudentCurrentStatus(StudentId=row.StudentId, CurrentStatus=status, LastPunchType=status,
                                LastPunchTime=when or NOW - timedelta(hours=25)))
    db.flush()
    return row


def punch(db, row, kind, when, status='VALID'):
    sid = db.query(RawPunch).count() + 1
    raw = RawPunch(SourceDeviceLogId=sid, UserId=row.UserId, DeviceId=101, LogDate=when)
    db.add(raw)
    db.flush()
    db.add(ProcessedPunch(RawPunchId=raw.RawPunchId, StudentId=row.StudentId, UserId=row.UserId,
                          DeviceId=101, LogDate=when, PunchType=kind, ProcessingStatus=status,
                          HostelIdSnapshot=row.HostelId))
    db.flush()


def test_current_lists_scope_hostel_export_and_read_purity(monkeypatch):
    monkeypatch.setattr(service, 'local_now', lambda: NOW)
    with TargetSessionLocal() as db:
        student(db, 'A_IN', status='IN')
        student(db, 'A_OUT')
        student(db, 'B_OUT', hostel=2)
        student(db, 'INACTIVE', registration='INACTIVE')
        student(db, 'PENDING', registration='PENDING_ENROLLMENT')
        h, w, it, mg = headers(db), headers(db, 'warden_a'), headers(db, 'it_admin'), headers(db, 'mgmt_test')
        db.commit()
        marker = change_marker(db)
    with TestClient(app) as client:
        assert [row['UserId'] for row in client.get('/api/student-details?view=in', headers=h).json()] == ['A_IN']
        for auth in (h, it, mg):
            assert {row['UserId'] for row in client.get('/api/student-details?view=out', headers=auth).json()} == {'A_OUT', 'B_OUT'}
        assert {row['UserId'] for row in client.get('/api/student-details?view=out&hostel_id=2', headers=h).json()} == {'B_OUT'}
        assert [row['UserId'] for row in client.get('/api/student-details?view=out', headers=w).json()] == ['A_OUT']
        for path in ('/api/student-details?view=out&hostel_id=2', '/api/student-details/export?view=out&hostel_id=2'):
            assert client.get(path, headers=w).status_code == 403
        exported = client.get('/api/student-details/export?view=out&search=A_OUT', headers=w)
        assert exported.status_code == 200 and exported.headers['x-total-count'] == '1'
        book = load_workbook(BytesIO(exported.content), read_only=True)
        rows = list(book.active.values)
        assert rows[3][0] == 'A_OUT' and len(rows) == 4
        book.close()
        assert client.get('/api/student-details').status_code == 401
        assert client.get('/api/student-details?view=wrong', headers=h).status_code == 422
        assert client.get('/api/student-details?limit=501', headers=h).status_code == 422
    with TargetSessionLocal() as db:
        assert marker == change_marker(db)
        assert db.query(StudentCurrentStatus).count() == 5


def test_five_threshold_separate_directions_and_india_day(monkeypatch):
    monkeypatch.setattr(service, 'local_now', lambda: NOW)
    day = datetime(2026, 10, 10)
    with TargetSessionLocal() as db:
        a, b, c, foreign = student(db, 'FIVE_IN'), student(db, 'FIVE_OUT'), student(db, 'COMBINED'), student(db, 'FOREIGN', hostel=2)
        for row, count_in, count_out in ((a, 5, 1), (b, 1, 6), (c, 3, 2), (foreign, 5, 0)):
            for index in range(count_in):
                punch(db, row, 'IN', day + timedelta(minutes=index))
            for index in range(count_out):
                punch(db, row, 'OUT', day + timedelta(hours=1, minutes=index))
        for when, kind, status in [(day-timedelta(seconds=1), 'IN', 'VALID'),
                                    (day+timedelta(days=1), 'IN', 'VALID'),
                                    (day+timedelta(hours=2), 'OUT', 'DUPLICATE'),
                                    (day+timedelta(hours=3), 'IN', 'INVALID_DEVICE'),
                                    (day+timedelta(hours=4), 'REGISTRATION', 'VALID')]:
            punch(db, c, kind, when, status)
        # Manual events do not become biometric punch counts.
        db.add(AttendanceOverride(StudentId=c.StudentId, PreviousStatus='IN', CurrentStatus='OUT',
                                 EffectiveAt=NOW, Reason='Test', CreatedBy='admin', HostelIdSnapshot=1))
        h, w = headers(db), headers(db, 'warden_a')
        db.commit()
    with TestClient(app) as client:
        result = client.get('/api/student-details?view=frequent', headers=h).json()
        assert {row['UserId'] for row in result} == {'FIVE_IN', 'FIVE_OUT', 'FOREIGN'}
        a_result = next(row for row in result if row['UserId'] == 'FIVE_IN')
        assert a_result['TodayInCount'] == 5 and a_result['TodayOutCount'] == 1
        assert {row['UserId'] for row in client.get('/api/student-details?view=frequent&direction=in', headers=w).json()} == {'FIVE_IN'}
        assert {row['UserId'] for row in client.get('/api/student-details?view=frequent&direction=out', headers=w).json()} == {'FIVE_OUT'}
        exported = client.get('/api/student-details/export?view=frequent&direction=out', headers=w)
        book = load_workbook(BytesIO(exported.content), read_only=True)
        assert [row[0] for row in list(book.active.values)[3:]] == ['FIVE_OUT']
        book.close()
    monkeypatch.setattr(service, 'local_now', lambda: day + timedelta(days=2))
    with TestClient(app) as client:
        assert client.get('/api/student-details?view=frequent', headers=h).json() == []


def test_out_duration_manual_reset_strict_boundary_and_unknown(monkeypatch):
    monkeypatch.setattr(service, 'local_now', lambda: NOW)
    with TargetSessionLocal() as db:
        student(db, 'OLD_OUT', when=NOW-timedelta(hours=25))
        student(db, 'EXACT24', when=NOW-timedelta(hours=24))
        student(db, 'CURRENT_IN', status='IN', when=NOW-timedelta(hours=30))
        unknown = student(db, 'UNKNOWN')
        curr = db.get(StudentCurrentStatus, unknown.StudentId)
        curr.LastPunchTime = None
        curr.LastPunchType = None
        recent = student(db, 'MANUAL_RECENT', when=NOW-timedelta(hours=30))
        old = student(db, 'MANUAL_OLD', status='IN')
        for row, hours in ((recent, 2), (old, 26)):
            override = AttendanceOverride(StudentId=row.StudentId, PreviousStatus='IN', CurrentStatus='OUT',
                EffectiveAt=NOW-timedelta(hours=hours), Reason='Test manual OUT', CreatedBy='admin', HostelIdSnapshot=1)
            db.add(override)
            db.flush()
            apply_override(db, row, override)
        student(db, 'OTHER_HOSTEL', hostel=2, when=NOW-timedelta(hours=28))
        h, w = headers(db), headers(db, 'warden_a')
        db.commit()
    with TestClient(app) as client:
        rows = client.get('/api/student-details?view=out_24h', headers=w).json()
        assert {row['UserId'] for row in rows} == {'OLD_OUT', 'MANUAL_OLD'}
        manual = next(row for row in rows if row['UserId'] == 'MANUAL_OLD')
        assert manual['OutHours'] == 26 and manual['OutSince'].endswith('+05:30')
        assert {row['UserId'] for row in client.get('/api/student-details?view=out_24h', headers=h).json()} == {'OLD_OUT', 'MANUAL_OLD', 'OTHER_HOSTEL'}


def test_export_all_pages_and_safe_cells(monkeypatch):
    monkeypatch.setattr(service, 'local_now', lambda: NOW)
    with TargetSessionLocal() as db:
        for i in range(105):
            row = student(db, f'EXPORT{i:03}')
            row.StudentName = '=DANGEROUS()' if i == 0 else row.StudentName
        h = headers(db)
        db.commit()
    with TestClient(app) as client:
        first = client.get('/api/student-details?view=out&search=EXPORT&limit=100', headers=h)
        last = client.get('/api/student-details?view=out&search=EXPORT&skip=100&limit=100', headers=h)
        assert first.headers['x-total-count'] == '105' and len(first.json()) == 100 and len(last.json()) == 5
        exported = client.get('/api/student-details/export?view=out&search=EXPORT', headers=h)
        assert exported.headers['x-total-count'] == '105'
        book = load_workbook(BytesIO(exported.content), read_only=True)
        rows = list(book.active.values)[3:]
        assert len(rows) == 105 and rows[0][1] == "'=DANGEROUS()"
        assert {row[0] for row in rows} == {row['UserId'] for row in first.json()+last.json()}
        book.close()


def test_warden_without_hostel_cannot_list_or_export(monkeypatch):
    monkeypatch.setattr(service, 'local_now', lambda: NOW)
    with TargetSessionLocal() as db:
        user = db.query(User).filter_by(Username='warden_a').one()
        user.HostelId = None
        h = headers(db, 'warden_a')
        db.commit()
    with TestClient(app) as client:
        for path in ('/api/student-details', '/api/student-details/export'):
            assert client.get(path, headers=h).status_code == 403
