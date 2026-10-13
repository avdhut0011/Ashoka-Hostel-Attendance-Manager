from typing import Literal, Optional
from tempfile import SpooledTemporaryFile
from fastapi import APIRouter, Depends, Query, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from backend.app.core.database import get_db
from backend.app.core.security import get_current_user
from backend.app.core.time import business_day
from backend.app.models.master import Student
from backend.app.models.security import User
from backend.app.services.student_details import detailed_query, detail_row
from backend.app.api.reports import safe_cell
from backend.app.core.report_limit import bounded_export
from backend.app.services.daily_counters import daily_counters

router = APIRouter(prefix='/student-details', tags=['Detailed student lists'])
ListView = Literal['in', 'out', 'frequent', 'out_24h']
Direction = Literal['either', 'in', 'out']


@router.get('')
def list_details(response: Response, view: ListView = 'in', hostel_id: Optional[int] = Query(None, ge=1),
                 search: Optional[str] = Query(None, max_length=150), direction: Direction = 'either',
                 skip: int = Query(0, ge=0), limit: int = Query(100, ge=1, le=500),
                 db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    query, now = detailed_query(db, current_user, view, hostel_id, search, direction, include_counts=view == 'frequent')
    response.headers['X-Total-Count'] = str(query.count())
    rows = list(query.order_by(Student.UserId, Student.StudentId).offset(skip).limit(limit))
    counts = daily_counters(db, [row.StudentId for row in rows], business_day(now)) if view != 'frequent' else {}
    result = [detail_row(row, now) for row in rows]
    if counts:
        for row in result:
            row.update({key: counts[row['StudentId']][key] for key in ('TodayInCount', 'TodayOutCount')})
    return result


@router.get('/export')
@bounded_export
def export_details(view: ListView = 'in', hostel_id: Optional[int] = Query(None, ge=1),
                   search: Optional[str] = Query(None, max_length=150), direction: Direction = 'either',
                   db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    query, now = detailed_query(db, current_user, view, hostel_id, search, direction)
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode='w+b')
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('Student List')
    sheet.append(['List', view, 'Generated (Asia/Kolkata)', now.isoformat(sep=' '), 'India attendance date', str(business_day(now))])
    sheet.append(['Hostel filter', current_user.HostelId if current_user.Role == 'WARDEN' else hostel_id or 'All', 'Search', safe_cell(search or ''), 'Punch direction', direction])
    columns = ['User ID', 'Student Name', 'Hostel', 'Room', 'Current Status', 'Today Valid IN', 'Today Valid OUT',
               'Last Event (Asia/Kolkata)', 'OUT Since (Asia/Kolkata)', 'Hours OUT']
    sheet.append(columns)
    count, sheet_rows, sheet_number = 0, 3, 1
    try:
        for row in query.order_by(Student.UserId, Student.StudentId).yield_per(500):
            if sheet_rows >= 1048576:
                sheet_number += 1
                sheet = workbook.create_sheet(f'Student List {sheet_number}')
                sheet.append(columns)
                sheet_rows = 1
            # Keep wall-clock values without offsets so Excel displays the same India time.
            sheet.append([safe_cell(row.UserId), safe_cell(row.StudentName or ''), safe_cell(row.HostelName or 'Unassigned'),
                          safe_cell(row.RoomNo or ''), row.CurrentStatus, int(row.TodayInCount), int(row.TodayOutCount),
                          row.LastPunchTime.isoformat(sep=' ') if row.LastPunchTime else '',
                          row.OutSince.isoformat(sep=' ') if row.OutSince else '',
                          round(max(0, (now - row.OutSince).total_seconds()/3600), 2) if row.OutSince else None])
            count += 1
            sheet_rows += 1
        workbook.save(output)
        output.seek(0)
    except Exception:
        output.close()
        raise
    finally:
        workbook.close()
    def chunks():
        try:
            while block := output.read(65536):
                yield block
        finally:
            output.close()
    return StreamingResponse(chunks(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'X-Total-Count': str(count), 'Content-Disposition': f'attachment; filename="students_{view}_{business_day(now)}.xlsx"'})
