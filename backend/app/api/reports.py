from datetime import date
from typing import List, Optional, Literal
import csv
from tempfile import SpooledTemporaryFile
from fastapi import APIRouter, Depends, Query, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from backend.app.core.database import get_db
from backend.app.core.security import get_current_user
from backend.app.core.time import business_day, local_now
from backend.app.models.attendance import ProcessedPunch
from backend.app.models.security import User
from backend.app.schemas.schemas import ProcessedPunchResponse
from backend.app.services.reporting import punch_query, punch_response, override_query, punch_count, punch_page
from backend.app.core.report_limit import bounded_export

router = APIRouter(prefix='/reports', tags=['Reports'])

@router.get('/today', response_model=List[ProcessedPunchResponse])
def get_today_report(response: Response, skip: int = Query(0, ge=0), limit: int = Query(100, ge=1, le=500), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    today = business_day(local_now())
    query = punch_query(db, current_user, today, today)
    response.headers['X-Total-Count'] = str(punch_count(query))
    return [punch_response(row, db) for row in punch_page(query, skip, limit)]

@router.get('/date-range', response_model=List[ProcessedPunchResponse])
def get_date_range_report(response: Response, start_date: date, end_date: date, hostel_id: Optional[int] = None,
    status: Optional[str] = None, user_id: Optional[str] = None, punch_type: Optional[str] = None,
    skip: int = Query(0, ge=0), limit: int = Query(100, ge=1, le=500), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    query = punch_query(db, current_user, start_date, end_date, hostel_id, user_id, punch_type, status)
    response.headers['X-Total-Count'] = str(punch_count(query))
    return [punch_response(row, db) for row in punch_page(query, skip, limit)]


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + value
    return value


def export_rows(query, overrides):
    from backend.app.services.verified_history import historical_room
    yield ['Event ID', 'User ID', 'Student Name', 'Hostel', 'Room No', 'Device ID', 'Event Date/Time (Asia/Kolkata)',
           'Punch Type', 'Processing Status', 'Reason', 'Historical Assignment Known', 'Changed By', 'Historical Room Source']
    for p, name, hostel, device, room in query.order_by(ProcessedPunch.LogDate.asc(), ProcessedPunch.PunchId.asc()).yield_per(500):
        historical, evidence = historical_room(query.session, p)
        yield [p.PunchId, p.UserId, name or '', hostel or '', historical or '',
            p.DeviceId, p.LogDate.isoformat(sep=' '), p.PunchType or '', p.ProcessingStatus, p.Reason or '',
            'YES' if p.HostelIdSnapshot is not None else 'NO (historical assignment unknown)', '', evidence]
    for o, uid, name, hostel in overrides.order_by('EffectiveAt').yield_per(500):
        yield [f'OVERRIDE-{o.OverrideId}', uid, name or '', hostel or '', o.RoomNoSnapshot or '', '',
            o.EffectiveAt.isoformat(sep=' '), o.CurrentStatus, 'MANUAL_OVERRIDE', o.Reason, 'YES', o.CreatedBy, 'OVERRIDE_SNAPSHOT']

@router.get('/export')
@bounded_export
def export_attendance_excel(start_date: Optional[date] = None, end_date: Optional[date] = None,
    hostel_id: Optional[int] = None, user_id: Optional[str] = None, punch_type: Optional[str] = None,
    status: Optional[str] = None, export_format: Literal['excel', 'csv'] = 'excel',
    include_overrides: bool = True, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    filters = dict(start=start_date, end=end_date, hostel_id=hostel_id, user_id=user_id, punch_type=punch_type, status=status)
    query = punch_query(db, current_user, **filters)
    overrides = override_query(db, current_user, **filters)
    if not include_overrides:
        from backend.app.models.attendance import AttendanceOverride
        overrides = overrides.filter(AttendanceOverride.OverrideId < 0)
    # Complete export before response: safe across FastAPI dependency lifetime versions.
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode='w+b')
    try:
        count = 0
        if export_format == 'csv':
            import io
            text = io.TextIOWrapper(output, encoding='utf-8-sig', newline='')
            writer = csv.writer(text)
            for row in export_rows(query, overrides):
                writer.writerow([safe_cell(v) for v in row])
                count += 1
            text.flush()
            text.detach()
            media, extension = 'text/csv; charset=utf-8', 'csv'
        else:
            workbook = Workbook(write_only=True)
            sheet_number = 1
            sheet = workbook.create_sheet('Attendance Report')
            header = None
            sheet_rows = 0
            for row in export_rows(query, overrides):
                if header is None:
                    header = row
                if sheet_rows >= 1048576:
                    sheet_number += 1
                    sheet = workbook.create_sheet(f'Attendance {sheet_number}')
                    sheet.append(header)
                    sheet_rows = 1
                sheet.append([safe_cell(v) for v in row])
                sheet_rows += 1
                count += 1
            workbook.save(output)
            workbook.close()
            media, extension = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'xlsx'
        output.seek(0)
    except Exception:
        output.close()
        raise
    def chunks():
        try:
            while chunk := output.read(65536):
                yield chunk
        finally:
            output.close()
    return StreamingResponse(chunks(), media_type=media, headers={
        'Content-Disposition': f'attachment; filename=attendance_report.{extension}',
        'X-Total-Count': str(max(0, count - 1))})
